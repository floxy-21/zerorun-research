"""Create and structurally verify the journal's editable Word highlights file.

Run with the configured bundled Python. Render the resulting DOCX separately
with the canonical document renderer and visually inspect every output page.
"""
from __future__ import annotations

import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET

ROOT = Path(__file__).resolve().parents[2]
HERE = ROOT / "research/softwarex"
TITLE = "ZeroRun highlights"
NS = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
      "dc": "http://purl.org/dc/elements/1.1/"}


def sha(path):
    return hashlib.sha256(Path(path).read_bytes()).hexdigest()


def highlights(path):
    lines = Path(path).read_text(encoding="utf-8").splitlines()
    if len(lines) != 5 or any(not line or line != line.strip() or len(line) > 85 for line in lines):
        raise ValueError("the unchanged text must contain exactly five nonempty highlights of at most 85 characters")
    return lines


def validate_docx(path, text_path):
    lines = highlights(text_path)
    with zipfile.ZipFile(path) as archive:
        if archive.testzip() is not None:
            raise ValueError("DOCX archive integrity failed")
        document = ET.fromstring(archive.read("word/document.xml"))
        numbering = ET.fromstring(archive.read("word/numbering.xml"))
        styles = ET.fromstring(archive.read("word/styles.xml"))
        core = ET.fromstring(archive.read("docProps/core.xml"))
        if core.findtext("dc:creator", namespaces=NS) != "Jishan Kapoor":
            raise ValueError("Word author metadata differs")
        paragraphs = document.findall("w:body/w:p", NS)
        texts = ["".join(p.itertext()) for p in paragraphs]
        if texts != [TITLE, *lines]:
            raise ValueError("Word paragraphs differ from the unchanged five source highlights")
        title_style = paragraphs[0].find("w:pPr/w:pStyle", NS)
        if title_style is None or title_style.get("{" + NS["w"] + "}val") != "Title":
            raise ValueError("document title does not use Word's Title style")
        for name in ("Normal", "Title"):
            style = styles.find("w:style[@w:styleId='" + name + "']", NS)
            if style is None or style.find("w:pPr/w:pBdr", NS) is not None:
                raise ValueError("unexpected paragraph border in document styles")
            fonts = style.find("w:rPr/w:rFonts", NS)
            if fonts is None or any("theme" in key.lower() for key in fonts.attrib):
                raise ValueError("theme font overrides must not override the explicit font")
            if any(fonts.get("{" + NS["w"] + "}" + key) != "Times New Roman" for key in ("ascii", "hAnsi")):
                raise ValueError("document font is not the requested Times New Roman")
            size = style.find("w:rPr/w:sz", NS)
            color = style.find("w:rPr/w:color", NS)
            if size is None or size.get("{" + NS["w"] + "}val") != "22" or color is None or color.get("{" + NS["w"] + "}val") != "000000":
                raise ValueError("document styles must use black 11-point text")
        for paragraph in paragraphs[1:]:
            num = paragraph.find("w:pPr/w:numPr/w:numId", NS)
            if num is None:
                raise ValueError("highlight is not a real Word list item")
            num_id = num.get("{" + NS["w"] + "}val")
            instance = numbering.find("w:num[@w:numId='" + num_id + "']/w:abstractNumId", NS)
            if instance is None:
                raise ValueError("Word list definition is missing")
            abstract = instance.get("{" + NS["w"] + "}val")
            form = numbering.find("w:abstractNum[@w:abstractNumId='" + abstract + "']/w:lvl/w:numFmt", NS)
            if form is None or form.get("{" + NS["w"] + "}val") != "bullet":
                raise ValueError("Word list is not a bullet list")
        sections = document.findall("w:body/w:sectPr", NS)
        if len(sections) != 1:
            raise ValueError("highlights must have one document section")
        page = sections[0].find("w:pgSz", NS)
        if page is None or page.get("{" + NS["w"] + "}w") != "12240" or page.get("{" + NS["w"] + "}h") != "15840":
            raise ValueError("Word page is not US Letter portrait")
        if sections[0].find("w:headerReference", NS) is not None or sections[0].find("w:footerReference", NS) is not None:
            raise ValueError("unexpected header or footer")
        if any(name.startswith(("word/comments", "word/header", "word/footer")) for name in archive.namelist()):
            raise ValueError("unexpected comments/header/footer part")
        if document.findall(".//w:ins", NS) or document.findall(".//w:del", NS):
            raise ValueError("tracked changes must not be present")
    return {"docx_sha256": sha(path), "source_text_sha256": sha(text_path),
            "author": "Jishan Kapoor", "highlight_count": 5,
            "highlight_character_counts": [len(line) for line in lines],
            "exact_source_text": True, "true_word_bullets": True,
            "letter_portrait": True, "header_footer_absent": True,
            "black_11pt_text_no_borders": True,
            "structural_checks_passed": True}


def create(output=HERE / "HIGHLIGHTS.docx", text_path=HERE / "HIGHLIGHTS.txt"):
    from docx import Document
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Inches, Pt, RGBColor

    output, text_path = Path(output), Path(text_path)
    if output.exists():
        raise ValueError("preserve the existing Word artifact before creating a new version")
    lines = highlights(text_path)
    source_before = sha(text_path)
    document = Document()
    properties = document.core_properties
    properties.author = "Jishan Kapoor"
    properties.last_modified_by = "Jishan Kapoor"
    properties.title = TITLE
    properties.subject = "SoftwareX submission highlights"
    properties.comments = ""
    properties.keywords = "ZeroRun, SoftwareX, research software"
    properties.created = datetime(2026, 9, 7, tzinfo=timezone.utc)
    properties.modified = datetime(2026, 9, 7, tzinfo=timezone.utc)
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    for name in ("Normal", "Title"):
        style = document.styles[name]
        style.font.name = "Times New Roman"
        style.font.size = Pt(11)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.underline = False
        style._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), "Times New Roman")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        for key in list(style._element.rPr.rFonts.attrib):
            if "theme" in key.lower():
                del style._element.rPr.rFonts.attrib[key]
        paragraph_properties = style._element.get_or_add_pPr()
        for border in paragraph_properties.findall(qn("w:pBdr")):
            paragraph_properties.remove(border)
        style.paragraph_format.space_after = Pt(8)
        style.paragraph_format.line_spacing = 1.15
    document.styles["Title"].font.bold = True
    document.styles["Title"].paragraph_format.space_after = Pt(14)
    document.add_paragraph(TITLE, "Title")
    numbering = document.part.numbering_part.element
    abstract_ids = [int(node.get(qn("w:abstractNumId"))) for node in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))]
    abstract_id, num_id = max(abstract_ids, default=-1) + 1, max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    for tag, value in (("w:start", "1"), ("w:numFmt", "bullet"), ("w:lvlText", "•"), ("w:lvlJc", "left")):
        element = OxmlElement(tag)
        element.set(qn("w:val"), value)
        level.append(element)
    paragraph_properties = OxmlElement("w:pPr")
    indent = OxmlElement("w:ind")
    indent.set(qn("w:left"), "288")
    indent.set(qn("w:hanging"), "216")
    paragraph_properties.append(indent)
    level.append(paragraph_properties)
    abstract.append(level)
    numbering.append(abstract)
    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_reference = OxmlElement("w:abstractNumId")
    abstract_reference.set(qn("w:val"), str(abstract_id))
    num.append(abstract_reference)
    numbering.append(num)
    for text in lines:
        paragraph = document.add_paragraph(text)
        paragraph.paragraph_format.keep_together = True
        properties = paragraph._p.get_or_add_pPr()
        num_properties = OxmlElement("w:numPr")
        for tag, value in (("w:ilvl", "0"), ("w:numId", str(num_id))):
            child = OxmlElement(tag)
            child.set(qn("w:val"), value)
            num_properties.append(child)
        properties.append(num_properties)
    document.save(output)
    if sha(text_path) != source_before:
        raise ValueError("highlight source text changed while creating Word artifact")
    return validate_docx(output, text_path)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=HERE / "HIGHLIGHTS.docx")
    parser.add_argument("--source", type=Path, default=HERE / "HIGHLIGHTS.txt")
    parser.add_argument("--check", action="store_true")
    args = parser.parse_args()
    result = validate_docx(args.output, args.source) if args.check else create(args.output, args.source)
    print(json.dumps(result, sort_keys=True))


if __name__ == "__main__":
    main()
